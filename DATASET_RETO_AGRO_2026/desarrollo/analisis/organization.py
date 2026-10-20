#Diccionario, clases y exportación de datos organizados

from __future__ import annotations
import csv
from collections import defaultdict
from pathlib import Path
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from analisis.config import CLASS_NOTES, EVOLUTION_PATH, SPECTRAL_INDICES, TARGET
from analisis.data import _number

__all__ = [
    "_variable_dictionary",
    "_write_reorganized_csv",
    "_write_workbook",
]

def _class_for_column(column: str) -> tuple[str, str, str]:
    source_column = column.removeprefix("HISTORICO_PROMEDIO__")
    upper = column.upper()
    if column == "ID_POLIGONO":
        return "Identificación y referencia", "sin unidad", "Clave primaria de parcela."
    if column in {"PARTICION", "FUENTE"} or upper.startswith(("FUENTE_", "SENSOR_")):
        return "Identificación y referencia", "categoría", "Metadato de partición o fuente/sensor."
    if "nubosidad" in source_column.lower():
        return "Calidad de observación", "%", "Porcentaje de nubosidad asociado con la captura."
    if column == "FECHA_CAPTURA":
        return "Tiempo de captura", "fecha", "Fecha de la observación satelital."
    if column in {"ANIO_CAPTURA", "MES_CAPTURA", "DIA_DEL_ANIO", "ANIO"}:
        return "Tiempo de captura", "año / mes / día", "Componente de calendario."
    if column == "AREA_HA":
        return "Identificación y referencia", "ha", "Área declarada en la tabla de parcelas."
    if column == TARGET:
        return "Variable objetivo", "t/ha", (
            "Rendimiento etiquetado en la tabla de asignación; la fuente no indica "
            "el ciclo del rendimiento. Solo se usa como objetivo, nunca como predictor."
        )
    if source_column.startswith("TOPOGRAFIA__elevacion"):
        return "Topografía", "m s.n.m.", "Media zonal del ráster INEGI CEM4."
    if source_column.startswith("TOPOGRAFIA__pendiente"):
        return "Topografía", "grados", "Media zonal del ráster derivado del CEM4."
    if source_column.startswith("CLIMA__precipitacion_mm_"):
        month = int(source_column.rsplit("_", 1)[1])
        return (
            "Precipitación",
            "mm/mes",
            f"Acumulado mensual CHIRPS; mes {month:02d}.",
        )
    if source_column.startswith("CLIMA__tmin_c_") or source_column.startswith("CLIMA__tmax_c_"):
        month = int(source_column.rsplit("_", 1)[1])
        variable = "mínima" if "__tmin_" in source_column else "máxima"
        return (
            "Temperatura",
            "°C",
            f"Promedio mensual CHIRTS-ERA5; temperatura {variable}, mes {month:02d}.",
        )

    spectral_column = source_column
    for prefix in ("HISTORICO_PROMEDIO__", "BASICO__", "PRO__"):
        spectral_column = spectral_column.removeprefix(prefix)
    pieces = spectral_column.split("_")
    index_name = pieces[0].lower() if pieces else ""
    if index_name in SPECTRAL_INDICES:
        description, variable_class = SPECTRAL_INDICES[index_name]
        suffix = "_".join(pieces[1:]).lower()
        statistic = {
            "promedio": "valor promedio por parcela y captura",
            "std": "desviación estándar espacial por parcela y captura",
            "max": "valor máximo por parcela y captura",
            "min": "valor mínimo por parcela y captura",
        }.get(suffix, "")
        if variable_class == "Índices espectrales sin expansión":
            note = "Los metadatos no respaldan una expansión completa de la abreviatura."
        else:
            note = "La descripción es conceptual; no implica fórmula ni causalidad."
        if statistic:
            note = f"{note} Estadístico: {statistic}."
        return variable_class, "sin unidad física declarada", f"{description}. {note}"
    if source_column.startswith(("BASICO__", "PRO__")):
        return (
            "Índices espectrales sin expansión",
            "sin unidad física declarada",
            "Conservar identificación original; definición no localizada.",
        )
    return "Identificación y referencia", "sin unidad", "Variable conservada con su nombre original."

def _descriptive_name(column: str) -> str:
    variable_class, _, note = _class_for_column(column)
    del variable_class, note
    source_column = column.removeprefix("HISTORICO_PROMEDIO__")
    is_historical_summary = column.startswith("HISTORICO_PROMEDIO__")
    if column == "ID_POLIGONO":
        base = "Identificador de parcela"
    elif column == "PARTICION":
        base = "Partición del conjunto"
    elif column == "FUENTE":
        base = "Fuente de la observación"
    elif column == "FECHA_CAPTURA":
        base = "Fecha de captura"
    elif column == "FUENTE" or column.startswith("FUENTE_"):
        base = "Fuente de la observación"
    elif column.startswith("SENSOR_"):
        base = "Sensor de observación"
    elif "nubosidad" in source_column.lower():
        base = "Porcentaje de nubosidad"
    elif column == TARGET:
        base = "Rendimiento etiquetado (t/ha)"
    elif column == "AREA_HA":
        base = "Área de parcela"
    elif source_column.startswith("TOPOGRAFIA__elevacion"):
        base = "Elevación media del terreno"
    elif source_column.startswith("TOPOGRAFIA__pendiente"):
        base = "Pendiente media del terreno"
    elif source_column.startswith("CLIMA__precipitacion_mm_"):
        month = int(source_column.rsplit("_", 1)[1])
        base = f"Precipitación mensual acumulada de {_month_name(month)}"
    elif source_column.startswith(("CLIMA__tmin_c_", "CLIMA__tmax_c_")):
        month = int(source_column.rsplit("_", 1)[1])
        variable = "mínima" if "__tmin_" in source_column else "máxima"
        base = f"Temperatura {variable} promedio de {_month_name(month)}"
    elif column in {"ANIO_CAPTURA", "MES_CAPTURA", "DIA_DEL_ANIO"}:
        base = {
            "ANIO_CAPTURA": "Año de captura",
            "MES_CAPTURA": "Mes de captura",
            "DIA_DEL_ANIO": "Día del año de captura",
        }[column]
    else:
        spectral_column = source_column
        for prefix in ("BASICO__", "PRO__"):
            spectral_column = spectral_column.removeprefix(prefix)
        pieces = spectral_column.split("_")
        index_name = pieces[0].lower()
        description = SPECTRAL_INDICES.get(
            index_name, (f"Variable espectral {index_name.upper()}", "")
        )[0]
        suffix = "_".join(pieces[1:]).lower()
        statistic = {
            "promedio": "valor promedio",
            "std": "desviación estándar",
            "max": "valor máximo",
            "min": "valor mínimo",
        }.get(suffix)
        base = description if statistic is None else f"{description}, {statistic}"
        if is_historical_summary:
            base = f"Promedio histórico por parcela de {base.lower()}"
        elif column.startswith("PRO__"):
            base = f"{base} (Dataset PRO)"
        elif column.startswith("BASICO__"):
            base = f"{base} (Dataset Básico)"
    if is_historical_summary and not base.lower().startswith(
        "promedio histórico por parcela de "
    ):
        base = f"Promedio histórico por parcela de {base.lower()}"
    return f"{base} ({column})"

def _month_name(month: int) -> str:
    return (
        "enero febrero marzo abril mayo junio julio agosto septiembre "
        "octubre noviembre diciembre"
    ).split()[month - 1]

def _class_order(class_name: str) -> int:
    order = [
        "Identificación y referencia",
        "Tiempo de captura",
        "Índices de vegetación",
        "Estructura del cultivo",
        "Agua y estrés hídrico",
        "Residuos y labranza",
        "Índices espectrales sin expansión",
        "Precipitación",
        "Temperatura",
        "Topografía",
        "Calidad de observación",
        "Variable objetivo",
    ]
    return order.index(class_name) if class_name in order else len(order)



def _variable_dictionary(columns: list[str]) -> list[dict[str, str]]:
    """Describe variables con nombre legible, clase y respaldo disponible."""
    records = []
    for column in columns:
        variable_class, unit, note = _class_for_column(column)
        records.append(
            {
                "identificador_original": column,
                "nombre_completo_con_identificador": _descriptive_name(column),
                "clase": variable_class,
                "unidad": unit,
                "procedencia_y_nota": note,
            }
        )
    return records

def _write_workbook(
    columns: list[str],
    dictionary: list[dict[str, str]],
    topography: dict[str, dict[str, float]],
    output_path: Path,
) -> None:
    """Exporta hojas de datos organizados, diccionario y descripción de clases."""
    all_columns = [
        name
        for name in columns
        if name != "RENDIMIENTO_T_HA"
    ] + [
        "RENDIMIENTO_T_HA",
        "TOPOGRAFIA__elevacion_media_m_snm",
        "TOPOGRAFIA__pendiente_media_grados",
    ]
    groups: dict[str, list[str]] = defaultdict(list)
    for column in all_columns:
        groups[_class_for_column(column)[0]].append(column)
    ordered_classes = sorted(groups, key=_class_order)
    ordered_columns = [
        column for group in ordered_classes for column in groups[group]
    ]
    group_by_column: dict[str, str] = {
        column: group for group in ordered_classes for column in groups[group]
    }

    workbook = Workbook(write_only=True)
    data_sheet = workbook.create_sheet("Datos organizados")
    data_sheet.freeze_panes = "A2"
    heading = []
    last_class = None
    for column in ordered_columns:
        group = group_by_column[column]
        if last_class is not None and group != last_class:
            heading.append(None)
        heading.append(_descriptive_name(column))
        last_class = group
    styled_header = []
    for value in heading:
        cell = WriteOnlyCell(data_sheet, value=value)
        if value is not None:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="234A67")
        styled_header.append(cell)
    data_sheet.append(styled_header)
    last_class = None
    source_to_position: dict[str, int] = {}
    class_boundaries: list[int] = []
    output_index = 0
    for column in ordered_columns:
        group = group_by_column[column]
        if last_class is not None and group != last_class:
            class_boundaries.append(output_index)
            output_index += 1
        source_to_position[column] = output_index
        output_index += 1
        last_class = group
    data_row_count = 0
    with EVOLUTION_PATH.open("r", encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source)
        for row in reader:
            data_row_count += 1
            polygon_id = (row.get("ID_POLIGONO") or "").strip()
            values = [None] * output_index
            for column, position in source_to_position.items():
                if column in row:
                    raw = row.get(column, "")
                elif column == "TOPOGRAFIA__elevacion_media_m_snm":
                    raw = topography.get(polygon_id, {}).get(column, "")
                elif column == "TOPOGRAFIA__pendiente_media_grados":
                    raw = topography.get(polygon_id, {}).get(column, "")
                else:
                    raw = ""
                if raw is None or raw == "":
                    values[position] = None
                elif column in {"ID_POLIGONO", "PARTICION", "FUENTE", "FECHA_CAPTURA"} or column.startswith(("FUENTE_", "SENSOR_")):
                    values[position] = raw
                else:
                    numeric = _number(str(raw))
                    values[position] = numeric if numeric is not None else raw
            data_sheet.append(values)
    data_sheet.auto_filter.ref = (
        f"A1:{get_column_letter(output_index)}{data_row_count + 1}"
    )

    dictionary_sheet = workbook.create_sheet("Diccionario de variables")
    dictionary_headers = list(dictionary[0]) if dictionary else []
    dictionary_sheet.append(dictionary_headers)
    for item in dictionary:
        dictionary_sheet.append([item[name] for name in dictionary_headers])
    class_sheet = workbook.create_sheet("Descripción de clases")
    class_sheet.append(["clase", "criterio de agrupación", "variables"])
    for group in ordered_classes:
        variables = [
            item["nombre_completo_con_identificador"]
            for item in dictionary
            if item["clase"] == group
        ]
        class_sheet.append(
            [group, CLASS_NOTES.get(group, ""), "; ".join(variables)]
        )
    workbook.save(output_path)

def _write_reorganized_csv(
    columns: list[str],
    topography: dict[str, dict[str, float]],
    output_path: Path,
) -> None:
    """Escribe los datos de presentación ordenados por clase, sin separadores."""
    export_columns = list(
        dict.fromkeys(
            [
                name
                for name in columns
                if name
            ]
            + [
                "TOPOGRAFIA__elevacion_media_m_snm",
                "TOPOGRAFIA__pendiente_media_grados",
            ]
        )
    )
    export_columns.sort(
        key=lambda name: (
            _class_order(_class_for_column(name)[0]),
            columns.index(name) if name in columns else 0,
        )
    )
    descriptive_headers = [_descriptive_name(name) for name in export_columns]
    if len(descriptive_headers) != len(set(descriptive_headers)):
        raise ValueError(
            "Los nombres descriptivos generarían encabezados duplicados en el CSV."
        )
    with EVOLUTION_PATH.open("r", encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source)
        with output_path.open("w", encoding="utf-8-sig", newline="") as output:
            writer = csv.DictWriter(output, fieldnames=descriptive_headers)
            writer.writeheader()
            for row in reader:
                polygon_id = (row.get("ID_POLIGONO") or "").strip()
                formatted = {
                    _descriptive_name(name): (
                        topography.get(polygon_id, {}).get(name, "")
                        if name.startswith("TOPOGRAFIA__")
                        else row.get(name, "")
                    )
                    for name in export_columns
                }
                writer.writerow(formatted)