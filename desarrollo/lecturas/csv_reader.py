import csv
from collections.abc import Iterable, Sequence
from pathlib import Path
from openpyxl import load_workbook
from config import OUTPUT_FOLDER, PROJECT_ROOT
from procesamiento.data_normalization import (
    QualityReport,
    clean_text,
    is_missing,
    normalize_number,
    normalize_text,
)

SOURCE_STEM = "ID_area_rendimiento_70_30_Reto_AgroCebada"

def _find_column(headers: Sequence[object], expected_name: str) -> int:
    matches = [
        index
        for index, header in enumerate(headers)
        if normalize_text(header).replace(" ", "_") == expected_name.casefold()
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Se esperaba exactamente una columna {expected_name}; "
            f"se encontraron {len(matches)}."
        )
    return matches[0]

def find_input_file(directory: Path = PROJECT_ROOT) -> Path:
    supported_extensions = (".csv", ".xlsx", ".xlsm")
    candidates = [
        directory / f"{SOURCE_STEM}{extension}"
        for extension in supported_extensions
        if (directory / f"{SOURCE_STEM}{extension}").is_file()
    ]
    if len(candidates) != 1:
        if candidates:
            names = ", ".join(path.name for path in candidates)
            raise FileNotFoundError(
                f"Hay más de un archivo para la tabla de zonas: {names}"
            )
        raise FileNotFoundError(
            f"No se encontró {SOURCE_STEM} con extensión CSV, XLSX o XLSM "
            f"en {directory}."
        )
    return candidates[0]

def filter_records(
    headers: Sequence[object] | None,
    rows: Iterable[tuple[int, Sequence[object]]],
    quality: QualityReport | None = None,
) -> dict[str, dict[str, object]]:
    #Valida filas y devuelve únicamente parcelas asignadas a predicción
    if headers is None:
        raise ValueError("El archivo de zonas no contiene una fila de encabezados.")

    header_names = [str(header).strip() if header is not None else "" for header in headers]
    id_index = _find_column(headers, "ID_POLIGONO")
    conjunto_index = _find_column(headers, "CONJUNTO")
    records: dict[str, dict[str, object]] = {}

    for row_number, row in rows:
        if not row or all(is_missing(value) for value in row):
            if quality is not None:
                quality.add(
                    "FILA_COMPLETAMENTE_VACIA", "fila_vacia", row_number
                )
            continue
        if len(row) != len(headers):
            if quality is not None:
                quality.add(
                    "FILA",
                    "cantidad_de_campos_invalida",
                    row_number,
                    f"{len(row)} valores; se esperaban {len(headers)}",
                )
            if len(row) <= max(id_index, conjunto_index):
                continue
            row = list(row[: len(headers)])
            row.extend([""] * (len(headers) - len(row)))
        if normalize_text(row[conjunto_index]) != "prediccion":
            continue

        polygon_id = str(row[id_index]).strip()
        if not polygon_id:
            raise ValueError(f"La fila {row_number} de predicción no tiene ID_POLIGONO.")
        if polygon_id in records:
            raise ValueError(
                f"ID_POLIGONO duplicado en predicciones: {polygon_id} "
                f"(fila {row_number})."
            )

        records[polygon_id] = {
            name: value
            for name, value in zip(header_names, row)
            if name
        }
    return records

def read_prediction_records(
    input_path: Path,
    quality: QualityReport | None = None,
) -> tuple[list[str], dict[str, dict[str, object]]]:
    """Lee la tabla de asignaciones desde CSV o la primera hoja de Excel."""
    if input_path.suffix.lower() == ".csv":
        with input_path.open("r", encoding="utf-8-sig", newline="") as csv_file:
            rows = csv.reader(csv_file)
            headers = next(rows, None)
            if headers is None:
                raise ValueError(f"El archivo está vacío: {input_path}")
            header_names = [str(header).strip() if header is not None else "" for header in headers]
            records = filter_records(headers, enumerate(rows, start=2), quality)
    else:
        workbook = load_workbook(input_path, read_only=True, data_only=True)
        try:
            rows = workbook.worksheets[0].iter_rows(values_only=True)
            headers = next(rows, None)
            if headers is None:
                raise ValueError(f"La primera hoja está vacía: {input_path}")
            header_names = [
                str(header).strip() if header is not None else ""
                for header in headers
            ]
            records = filter_records(headers, enumerate(rows, start=2), quality)
        finally:
            workbook.close()
    return [name for name in header_names if name], records

def run_initial_reader() -> Path:
    #Normaliza asignaciones y exporta las parcelas reservadas para predecir
    input_path = find_input_file()
    quality = QualityReport("prediccion")
    headers, predictions = read_prediction_records(input_path, quality)
    numeric_fields = {"AREA_HA", "RENDIMIENTO_T_HA"}
    output_rows: list[dict[str, object]] = []

    for row_number, record in enumerate(predictions.values(), start=2):
        normalized_record: dict[str, object] = {}
        for field in headers:
            value = record.get(field)
            if is_missing(value):
                quality.add(field or "ENCABEZADO_VACIO", "faltante", row_number)
                normalized_record[field] = ""
            elif field in numeric_fields:
                try:
                    normalized_record[field] = normalize_number(value)
                except ValueError:
                    quality.add(field, "formato_invalido", row_number, value)
                    normalized_record[field] = ""
            else:
                normalized_record[field] = clean_text(value)
        output_rows.append(normalized_record)

    output_path = OUTPUT_FOLDER / "parcvelas_prediccion.csv"
    write_csv(output_path, headers, output_rows)
    quality.write(OUTPUT_FOLDER / "calidad_prediccion.csv")

    print(f"Archivo de entrada: {input_path}")
    print(f"Filas con CONJUNTO = predicción: {len(output_rows)}")
    print(f"CSV generado: {output_path}")
    print(f"Reporte de calidad: {OUTPUT_FOLDER / 'calidad_prediccion.csv'}")
    return output_path

def write_csv(
    output_path: Path,
    headers: Sequence[str],
    rows: Iterable[dict[str, object]],
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8-sig", newline="") as output_file:
        writer = csv.DictWriter(output_file, fieldnames=headers, extrasaction="raise")
        writer.writeheader()
        writer.writerows(rows)